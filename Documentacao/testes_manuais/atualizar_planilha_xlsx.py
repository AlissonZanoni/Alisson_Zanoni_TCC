from pathlib import Path
from openpyxl import load_workbook

xlsx = Path(r"C:\Users\aliss_16gza5b\OneDrive\Documentos\Alisson_Zanoni_TCC\Documentacao\testes_manuais\planilha_testes_manuais_profissional.xlsx")

wb = load_workbook(xlsx)
ws = wb["Execucao_QA"]

updates = {
    "HP-02": {
        "descricao": 'Pagina inicial com exatamente quatro "Arrivals"',
        "tipo": "Funcional",
        "passos": "1. Abrir o navegador e acessar a URL. 2. Observar a secao de Arrivals na home.",
        "resultado": "A pagina inicial deve conter exatamente 4 produtos na secao Arrivals.",
    },
    "SH-10": {
        "descricao": "Exibicao dos produtos com promocao obrigatoria",
        "tipo": "Funcional",
        "passos": "1. Acessar Shop. 2. Observar a listagem de produtos.",
        "resultado": "Cada produto deve exibir: titulo + preco promocional completo (preco antigo riscado e preco atual) + botao para adicionar ao carrinho.",
    },
    "LG-08": {
        "descricao": "Checkbox Remember me marcado por padrao",
        "tipo": "Funcional",
        "passos": "1. Acessar My Account. 2. Verificar o formulario de login.",
        "resultado": "O checkbox Remember me deve estar visivel, disponivel para interacao e iniciar marcado por padrao.",
    },
    "RG-03": {
        "descricao": "E-mail vazio - botao Register deve estar desabilitado",
        "tipo": "Funcional",
        "resultado": "O botao Register deve permanecer desabilitado quando o campo de e-mail estiver vazio.",
    },
    "RG-04": {
        "descricao": "Senha vazia - botao Register deve estar desabilitado",
        "tipo": "Funcional",
        "resultado": "O botao Register deve permanecer desabilitado quando o campo de senha estiver vazio.",
    },
    "RG-05": {
        "descricao": "E-mail e senha vazios - botao Register deve estar desabilitado",
        "tipo": "Funcional",
        "resultado": "O botao Register deve permanecer desabilitado quando os campos de e-mail e senha estiverem vazios.",
    },
    "CA-12": {
        "descricao": "Validar limite de produtos no carrinho",
        "tipo": "Funcional",
        "resultado": "O carrinho deve respeitar o limite maximo de produtos definido pela regra de negocio.",
    },
}

for row in range(2, ws.max_row + 1):
    test_id = ws.cell(row=row, column=1).value
    if test_id in updates:
        upd = updates[test_id]
        if "descricao" in upd:
            ws.cell(row=row, column=4, value=upd["descricao"])
        if "tipo" in upd:
            ws.cell(row=row, column=5, value=upd["tipo"])
        if "passos" in upd:
            ws.cell(row=row, column=8, value=upd["passos"])
        if "resultado" in upd:
            ws.cell(row=row, column=9, value=upd["resultado"])

wb.save(xlsx)
print("Planilha atualizada com sucesso")
