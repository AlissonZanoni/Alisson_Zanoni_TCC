from pathlib import Path
from openpyxl import load_workbook

xlsx = Path(r"C:\Users\aliss_16gza5b\OneDrive\Documentos\Alisson_Zanoni_TCC\Documentacao\testes_manuais\planilha_testes_manuais_profissional.xlsx")
wb = load_workbook(xlsx)
ws = wb["Execucao_QA"]

for row in range(2, ws.max_row + 1):
    if ws.cell(row=row, column=1).value == "SH-10":
        ws.cell(row=row, column=4, value="Exibicao dos produtos com destaque de promocao na primeira fileira")
        ws.cell(row=row, column=8, value="1. Acessar Shop. 2. Observar a listagem de produtos e identificar produtos em promocao.")
        ws.cell(row=row, column=9, value="Cada produto deve conter titulo, preco e botao de adicionar ao carrinho. Produtos em promocao devem estar posicionados na primeira fileira da listagem.")
        ws.cell(row=row, column=5, value="Funcional")
        break

wb.save(xlsx)
print("SH-10 atualizado na planilha")
