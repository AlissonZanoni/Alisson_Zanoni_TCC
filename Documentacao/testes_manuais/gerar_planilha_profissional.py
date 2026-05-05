from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE_DIR = Path(__file__).resolve().parent
CSV_PATH = BASE_DIR / "planilha_testes_manuais.csv"
XLSX_PATH = BASE_DIR / "planilha_testes_manuais_profissional.xlsx"

STATUS_OPCOES = [
    "Não Executado",
    "Em Andamento",
    "Testado com Sucesso",
    "Testado com Falhas",
]


def ler_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    with path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.reader(f, delimiter=";")
        rows = list(reader)

    if not rows:
        raise ValueError("CSV vazio")

    header = rows[0]
    data = rows[1:]
    return header, data


def ajustar_larguras(ws) -> None:
    max_width = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            value = "" if cell.value is None else str(cell.value)
            max_width[cell.column_letter] = max(max_width.get(cell.column_letter, 0), len(value))

    for col, width in max_width.items():
        ws.column_dimensions[col].width = min(max(width + 2, 12), 55)


def criar_aba_execucao(wb: Workbook, header: list[str], data: list[list[str]]) -> tuple[object, int]:
    ws = wb.active
    ws.title = "Execucao_QA"

    ws.append(header)
    for row in data:
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    header_fill = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    table_ref = f"A1:M{ws.max_row}"
    table = Table(displayName="TabelaExecucaoQA", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    ajustar_larguras(ws)

    # Colunas maiores para leitura
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["H"].width = 58
    ws.column_dimensions["I"].width = 58
    ws.column_dimensions["K"].width = 34
    ws.column_dimensions["L"].width = 42

    status_col = header.index("Status") + 1

    # Lista de status pre-definida em aba auxiliar (evita problema de separador regional)
    ws_listas = wb.create_sheet("Listas")
    for idx, status in enumerate(STATUS_OPCOES, start=1):
        ws_listas.cell(row=idx, column=1, value=status)
    ws_listas.sheet_state = "hidden"

    status_validation = DataValidation(type="list", formula1="=Listas!$A$1:$A$4", allow_blank=False)
    ws.add_data_validation(status_validation)
    status_validation.add(f"{ws.cell(row=1, column=status_col).column_letter}2:{ws.cell(row=ws.max_row, column=status_col).column_letter}{ws.max_row}")

    # Cores por status
    status_range = f"{ws.cell(row=2, column=status_col).column_letter}2:{ws.cell(row=ws.max_row, column=status_col).column_letter}{ws.max_row}"

    ws.conditional_formatting.add(
        status_range,
        FormulaRule(formula=[f"${ws.cell(row=2, column=status_col).column_letter}2=\"Não Executado\""], fill=PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")),
    )
    ws.conditional_formatting.add(
        status_range,
        FormulaRule(formula=[f"${ws.cell(row=2, column=status_col).column_letter}2=\"Em Andamento\""], fill=PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")),
    )
    ws.conditional_formatting.add(
        status_range,
        FormulaRule(formula=[f"${ws.cell(row=2, column=status_col).column_letter}2=\"Testado com Sucesso\""], fill=PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")),
    )
    ws.conditional_formatting.add(
        status_range,
        FormulaRule(formula=[f"${ws.cell(row=2, column=status_col).column_letter}2=\"Testado com Falhas\""], fill=PatternFill(start_color="FECACA", end_color="FECACA", fill_type="solid")),
    )

    # Cores por secao (coluna B, faixa total da linha)
    secao_cores = {
        "Home Page": "E8F4FD",
        "Shop": "FFF4E5",
        "My Account - Login": "F3E8FF",
        "My Account - Registro": "FFE4E6",
        "My Account - Área Logada": "ECFDF5",
        "My Account - Area Logada": "ECFDF5",
        "My Account - Account Details": "EEF2FF",
        "My Account - Addresses": "F0F9FF",
        "Carrinho": "FEFCE8",
    }

    faixa_linhas = f"A2:M{ws.max_row}"
    for secao, cor in secao_cores.items():
        ws.conditional_formatting.add(
            faixa_linhas,
            FormulaRule(formula=[f"$B2=\"{secao}\""], fill=PatternFill(start_color=cor, end_color=cor, fill_type="solid")),
        )

    return ws, status_col


def criar_dashboard(wb: Workbook, ws_execucao, status_col: int) -> None:
    ws = wb.create_sheet("Dashboard")

    ws["A1"] = "Dashboard de Execucao QA"
    ws["A1"].font = Font(size=16, bold=True, color="0F4C81")

    ws["A3"] = "Status"
    ws["B3"] = "Quantidade"
    ws["A3"].font = Font(bold=True)
    ws["B3"].font = Font(bold=True)

    status_col_letter = ws_execucao.cell(row=1, column=status_col).column_letter
    last_row = ws_execucao.max_row

    for i, status in enumerate(STATUS_OPCOES, start=4):
        ws[f"A{i}"] = status
        # Usa TRIM para ignorar espacos extras acidentais digitados manualmente.
        ws[f"B{i}"] = (
            f"=SUMPRODUCT(--(TRIM(Execucao_QA!${status_col_letter}$2:${status_col_letter}${last_row})=TRIM(A{i})))"
        )

    ws["A9"] = "Total de cenarios"
    ws["B9"] = "=SUM(B4:B7)"
    ws["A10"] = "Concluidos"
    ws["B10"] = "=B6+B7"
    ws["A11"] = "Progresso (%)"
    ws["B11"] = "=IF(B9=0,0,B10/B9)"
    ws["B11"].number_format = "0.00%"

    ws["A13"] = "Tempo total de execucao (geral)"
    ws["B13"] = None
    ws["B13"].number_format = "[h]:mm:ss"
    ws["C13"] = "Preencher em hh:mm:ss"
    ws["C13"].font = Font(italic=True, color="6B7280")

    tempo_fill = PatternFill(start_color="FFF7CC", end_color="FFF7CC", fill_type="solid")
    ws["B13"].fill = tempo_fill
    ws["B13"].alignment = Alignment(horizontal="center")

    for cell in ["A9", "A10", "A11", "A13"]:
        ws[cell].font = Font(bold=True)

    pie = PieChart()
    labels = Reference(ws, min_col=1, min_row=4, max_row=7)
    data = Reference(ws, min_col=2, min_row=3, max_row=7)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.title = "Distribuicao de Status"
    pie.style = 10
    pie.height = 9
    pie.width = 14
    ws.add_chart(pie, "D3")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 28


def main() -> None:
    header, data = ler_csv(CSV_PATH)

    wb = Workbook()
    ws_execucao, status_col = criar_aba_execucao(wb, header, data)
    criar_dashboard(wb, ws_execucao, status_col)

    wb.save(XLSX_PATH)
    print(f"Planilha criada: {XLSX_PATH}")


if __name__ == "__main__":
    main()
